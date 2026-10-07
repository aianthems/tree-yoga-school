"""Import the highest published scores from WV Forestry's 2025 registers.

Usage: python scripts/import-west-virginia-champions.py SNAPSHOT_DIR
Requires openpyxl. SNAPSHOT_DIR contains flowering.xlsx, conifers.xlsx and
counties.zip (2025 Census national county gazetteer). Sources are documented
in lib/champion-trees.ts. Never treats all registered big trees as champions.
The workbooks have no general champion designation or co-champion rule:
selection is derived from published BTP, including exact ties only.
"""
from collections import defaultdict
from pathlib import Path
import csv
import hashlib
import io
import json
import re
import sys
import zipfile
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
COUNTY_ALIASES = {"Pendelton": "Pendleton", "Pochontas": "Pocahontas"}


def clean(value):
    return " ".join(str(value).split()) if value is not None else ""


def number(value):
    assert isinstance(value, (int, float)), repr(value)
    return value


def extract(snapshot):
    with zipfile.ZipFile(snapshot / "counties.zip") as archive:
        rows = csv.DictReader(io.StringIO(archive.read(archive.namelist()[0]).decode()), delimiter="|")
        geography = {row["NAME"].removesuffix(" County"): {
            "lat": float(row["INTPTLAT"]), "lng": float(row["INTPTLONG"])
        } for row in rows if row["USPS"] == "WV"}
    groups = defaultdict(list)
    audit = {"edition": "2025", "sourceFiles": {}, "excludedCategories": [], "selection": []}
    for kind, title in [("flowering", "Flowering"), ("conifers", "Conifers")]:
        filename = snapshot / f"{kind}.xlsx"
        audit["sourceFiles"][kind] = {"sha256": hashlib.sha256(filename.read_bytes()).hexdigest()}
        sheet = openpyxl.load_workbook(filename, data_only=True).active
        assert [sheet.cell(1, n).value.strip() for n in [1, 2, 7, 8, 9, 10, 12]] == [
            "Scientific Name", "Common Name", "BTP", "County", "Town", "Year Nominated", "Last Inspection"]
        count = 0
        for row_no, row in enumerate(sheet.iter_rows(min_row=2, max_col=13, values_only=True), start=2):
            if not row[0]:
                continue
            count += 1
            category = clean(row[1]).replace("Chineese Chestnut", "Chinese Chestnut")
            groups[category].append((kind, title, row_no, row))
        audit["sourceFiles"][kind]["records"] = count
    records = []
    for category, entries in groups.items():
        if category == "Chinese Chestnut":
            # The 185-point leader is called Chinese Chestnut but has Castanea
            # dentata (American chestnut). Exclude the whole category rather than
            # promote its runner-up or silently assign an identity.
            audit["excludedCategories"].append({"category": category,
                "reason": "Leading row 103 has conflicting common/scientific names; no successor inferred.",
                "rows": [entry[2] for entry in entries]})
            continue
        maximum = max(number(entry[3][6]) for entry in entries)
        leaders = [entry for entry in entries if entry[3][6] == maximum]
        audit["selection"].append({"category": category, "maximumPublishedPoints": maximum,
            "sourceRowsCompared": [f"{entry[0]}:{entry[2]}" for entry in entries],
            "selectedRows": [f"{entry[0]}:{entry[2]}" for entry in leaders]})
        for kind, title, row_no, row in leaders:
            warnings = []
            scientific = clean(row[0])
            if " - " in scientific:
                original = scientific
                scientific = scientific.split(" - ", 1)[0]
                warnings.append(f"The scientific-name cell reads ‘{original}’. The annotation is retained here; any national label is not independently confirmed as current.")
            circumference = row[3]
            if isinstance(circumference, str):
                match = re.match(r"^(\d+(?:\.\d+)?)\s*(?:@|MS)", circumference)
                assert match, (kind, row_no, circumference)
                warnings.append(f"Published circumference entry: ‘{circumference}’. @ indicates a nonstandard measurement height; MS indicates a fused multi-stem bole.")
                circumference = float(match[1])
            county_original = clean(row[7])
            county = COUNTY_ALIASES.get(county_original, county_original)
            assert county in geography, (kind, row_no, county)
            if county != county_original:
                warnings.append(f"The source spells the county ‘{county_original}’; the approximate marker uses {county} County. The published town is retained.")
            assert isinstance(row[11], (int, float)) and 1900 <= row[11] <= 2025
            records.append({"id": f"wv-2025-{kind}-{row_no}", "state": "WV", "sourceRow": row_no,
                "scientificName": scientific, "commonName": clean(row[1]),
                "town": clean(row[8]), "county": county, "sourceCounty": county_original,
                "mapPrecision": "county", "location": None, "measured": str(int(row[11])),
                "circumference": number(circumference), "height": number(row[4]),
                "crown": number(row[5]), "points": number(row[6]),
                "status": "Joint highest published score" if len(leaders) > 1 else "Highest published score",
                "sourceUrl": f"https://wvforestry.com/pdf/bigtree/2025-{title}-Common.xlsx",
                "sourceReviewNotes": warnings,
                "notes": f"2025 {title} register · row {row_no}. Last inspection year is reported by the source; it is not the workbook edition date."})
    audit["selectedRecords"] = len(records)
    audit["selectedCategories"] = len(audit["selection"])
    used = {record["county"]: geography[record["county"]] for record in records}
    return records, dict(sorted(used.items())), audit


if __name__ == "__main__":
    records, points, audit = extract(Path(sys.argv[1]))
    assert len(records) == 159 and len(audit["selection"]) == 157
    assert len({record["id"] for record in records}) == len(records)
    for filename, data in [("west-virginia-champion-trees.json", records),
                           ("west-virginia-county-points.json", points),
                           ("west-virginia-import-audit.json", audit)]:
        # One record per line keeps this data file compact and changes reviewable.
        text = "[\n" + ",\n".join("  " + json.dumps(row, ensure_ascii=False) for row in data) + "\n]\n" if isinstance(data, list) else json.dumps(data, ensure_ascii=False, indent=2) + "\n"
        (ROOT / "lib/data" / filename).write_text(text)
    print(json.dumps({"records": len(records), "categories": len(audit["selection"]), "counties": len(points)}))
